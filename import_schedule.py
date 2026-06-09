"""
2026_보건_스케쥴.xlsx 전체 시트 파싱
"""
import openpyxl
import pandas as pd

AGENCY_MAP = {
    "중대협": "중대협",
    "수원_대한협": "수원대한협",
    "대한협_수원": "수원대한협",
    "서울_대한협": "서울대한협",
    "대한협_서울": "서울대한협",
    "인천_대한협": "인천대한협",
    "대한협_인천": "인천대한협",
    "한안협": "한안협",
    "잡그레이드": "잡그레이드",
}

LOCATION_MAP = {
    "온오프": "동시송출",
    "오프": "오프라인",
    "오프라인": "오프라인",
    "줌": "줌",
    "zoom": "줌",
    "2층": "수원2층",
    "5층": "수원5층",
    "오프-2층": "수원2층",
    "수원-2층": "수원2층",
}

def _norm_agency(text):
    text = str(text).strip()
    return AGENCY_MAP.get(text, text)

def _norm_industry(text):
    text = str(text).strip()
    if "제조" in text: return "제조업"
    if "건설" in text: return "건설업"
    if "보건" in text: return "기타업"
    if "서비스" in text: return "기타업"
    return "기타업" if text in ("기타업", "기타", "기타신규", "기타_신규", "") else text

def _norm_location(text, memo=""):
    text = str(text).strip()
    loc = LOCATION_MAP.get(text, text)
    if memo:
        m = str(memo)
        if "제1강의실" in m or "1강의실" in m: return "인천1"
        if "제2강의실" in m or "2강의실" in m: return "인천2"
        if "2층" in m and "수원" not in loc: return "수원2층"
        if "5층" in m: return "수원5층"
        if "광교" in m: return "수원광교"
    return loc

def _norm_time(val):
    if val is None: return ""
    if hasattr(val, 'strftime'): return val.strftime("%H:%M")
    try:
        h = int(float(val))
        m = int(round((float(val) - h) * 60))
        return f"{h:02d}:{m:02d}"
    except: return str(val)

def _norm_subject(text):
    text = str(text).strip()
    from utils import detect_subject
    return detect_subject(text) or text

def _norm_target(text):
    text = str(text).strip()
    if "책임자" in text or "관리감독" in text: return "관리감독자"
    if "안전관리" in text: return "안전관리자"
    if "보건관리" in text: return "보건관리자"
    if "근로자" in text: return "근로자"
    return text or "관리감독자"

def parse_all_sheets(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    all_rows = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows: continue

        # 헤더 찾기
        header_idx = 0
        for i, row in enumerate(rows):
            if row[0] and "강의" in str(row[0]):
                header_idx = i
                break

        for row in rows[header_idx + 1:]:
            if not row or row[0] is None: continue

            # 날짜
            try:
                if hasattr(row[0], 'strftime'):
                    date_str = row[0].strftime("%Y-%m-%d")
                    weekday = ["월","화","수","목","금","토","일"][row[0].weekday()]
                else:
                    dt = pd.to_datetime(str(row[0]))
                    date_str = dt.strftime("%Y-%m-%d")
                    weekday = ["월","화","수","목","금","토","일"][dt.weekday()]
            except: continue

            start = _norm_time(row[2]) if len(row) > 2 else ""
            end   = _norm_time(row[3]) if len(row) > 3 else ""
            agency    = _norm_agency(row[4])   if len(row) > 4 and row[4] else ""
            subject   = _norm_subject(row[5])  if len(row) > 5 and row[5] else ""
            target    = _norm_target(row[6])   if len(row) > 6 and row[6] else "관리감독자"
            industry  = _norm_industry(row[7]) if len(row) > 7 and row[7] else "기타업"
            memo1     = str(row[13]).strip()   if len(row) > 13 and row[13] else ""
            memo2     = str(row[14]).strip()   if len(row) > 14 and row[14] else ""
            location  = _norm_location(str(row[8]).strip() if len(row) > 8 and row[8] else "", memo1)
            instructor = str(row[9]).strip()   if len(row) > 9 and row[9] else ""

            try: hours = int(float(row[10])) if len(row) > 10 and row[10] else 0
            except: hours = 0
            try: hourly_fee = int(float(row[11])) if len(row) > 11 and row[11] else 100000
            except: hourly_fee = 100000
            try: daily_fee = int(float(row[12])) if len(row) > 12 and row[12] else hours * hourly_fee
            except: daily_fee = hours * hourly_fee

            all_rows.append({
                "강의일시": date_str,
                "요일": weekday,
                "시작": start,
                "종료": end,
                "의뢰기관": agency,
                "과정명": subject,
                "대상자": target,
                "업종": industry,
                "방식/위치": location,
                "강사님": instructor,
                "시수": hours,
                "강의료(1시간)": hourly_fee,
                "강의료(1일)": daily_fee,
                "상태": "정상",
                "요청사항": memo1,
                "내부메모": memo2,
                "의뢰일": "",
                "의뢰인": "",
                "의뢰방법": "",
                "변경일자": "",
                "변경이력": "",
                "변경의뢰인": "",
                "증빙폴더": "",
            })

    df = pd.DataFrame(all_rows)
    df["_dt"] = pd.to_datetime(df["강의일시"], errors="coerce")
    df = df.sort_values("_dt", ascending=True).drop(columns=["_dt"]).reset_index(drop=True)
    return df
