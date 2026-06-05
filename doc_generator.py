"""
DOCX 의뢰서 자동 생성 모듈
카톡/엑셀 데이터 → 강의 의뢰서.docx
"""

import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def _add_section_title(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def _add_field(doc: Document, label: str, value: str):
    p = doc.add_paragraph()
    label_run = p.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.size = Pt(11)
    value_run = p.add_run(str(value) if value else "-")
    value_run.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)


def _add_divider(doc: Document):
    p = doc.add_paragraph("─" * 50)
    p.runs[0].font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    p.runs[0].font.size = Pt(9)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)


def generate_request_docx(row: dict, original_text: str = "") -> bytes:
    """
    강의 의뢰서 DOCX 생성
    row: DataFrame 한 행 (dict)
    original_text: 카톡/이메일 원문
    반환: bytes (파일 저장 또는 드라이브 업로드용)
    """
    doc = Document()

    # 여백 설정
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # 제목
    _add_title(doc, "강 의 의 뢰 서")
    doc.add_paragraph()

    # 기본 정보
    _add_section_title(doc, "■ 강의 정보")
    _add_divider(doc)
    _add_field(doc, "의뢰기관",   row.get("의뢰기관", ""))
    _add_field(doc, "과정명",     row.get("과정명", ""))
    _add_field(doc, "대상자",     row.get("대상자", ""))
    _add_field(doc, "업종",       row.get("업종", ""))
    _add_field(doc, "방식/위치",  row.get("방식/위치", ""))
    doc.add_paragraph()

    # 일정 정보
    _add_section_title(doc, "■ 일정")
    _add_divider(doc)
    _add_field(doc, "강의일시",   row.get("강의일시", ""))
    _add_field(doc, "요일",       row.get("요일", ""))
    _add_field(doc, "시작시간",   row.get("시작", ""))
    _add_field(doc, "종료시간",   row.get("종료", ""))
    _add_field(doc, "시수",       f"{row.get('시수', '')}시간")
    doc.add_paragraph()

    # 강사/의뢰 정보
    _add_section_title(doc, "■ 강사 및 의뢰")
    _add_divider(doc)
    _add_field(doc, "강사",       row.get("강사님", ""))
    _add_field(doc, "시간당 강의료", f"₩{row.get('강의료(1시간)', 0):,}")
    _add_field(doc, "1일 강의료",  f"₩{row.get('강의료(1일)', 0):,}")
    _add_field(doc, "의뢰인",     row.get("의뢰인", ""))
    _add_field(doc, "의뢰일",     row.get("의뢰일", ""))
    _add_field(doc, "의뢰방법",   row.get("의뢰방법", ""))
    doc.add_paragraph()

    # 메모
    memo = row.get("의뢰업체메모", "")
    if memo:
        _add_section_title(doc, "■ 의뢰업체 메모")
        _add_divider(doc)
        p = doc.add_paragraph(str(memo))
        p.font_size = Pt(10)
        doc.add_paragraph()

    # 원본 내용
    if original_text and original_text.strip():
        _add_section_title(doc, "■ 원본 요청 내용")
        _add_divider(doc)
        p = doc.add_paragraph(original_text.strip())
        for run in p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        doc.add_paragraph()

    # 생성 일시
    p = doc.add_paragraph(f"문서 생성: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # bytes 반환
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_docx_filename(row: dict) -> str:
    """파일명 생성: 20260605_대한협수원_응급처치대한1.docx"""
    date_str = str(row.get("강의일시", "")).replace("-", "")
    agency = str(row.get("의뢰기관", "")).replace(" ", "")
    subject = str(row.get("과정명", "")).replace(" ", "")
    return f"{date_str}_{agency}_{subject}.docx"
